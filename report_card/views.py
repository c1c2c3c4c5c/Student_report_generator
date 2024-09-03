import openpyxl
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect
from django.contrib.auth import get_user_model
from django.db.models import *
from .models import *
from .forms import *
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from django.core.exceptions import ObjectDoesNotExist
from django.template.loader import get_template
from xhtml2pdf import pisa
from io import BytesIO
from django.http import FileResponse
from django.shortcuts import get_object_or_404
from django.http import HttpResponse


User = get_user_model()



@login_required
def mark_input(request):
    if request.method == "POST":
        form = MarkForm(request.POST)
        if form.is_valid():
            mark = form.save(commit=False)
            mark.rank = (
                SubjectMarks.objects.filter(
                    subject=mark.subject, marks__gt=mark.marks
                ).count()
                + 1
            )
            mark.save()
            # Calculate overall rank
            total_marks = (
                SubjectMarks.objects.filter(student=mark.student).aggregate(
                    total=Sum("marks")
                )["total"]
                or 0
            )
            overall_rank = (
                Student.objects.annotate(total_marks=Sum("studentmarks__marks"))
                .filter(total_marks__gt=total_marks)
                .count()
                + 1
            )
            # Save overall rank to the student record
            mark.student.overall_rank = overall_rank
            mark.student.save()
            return redirect("mark_input_success")
    else:
        form = MarkForm()
    return render(request, "report_card/mark_input.html", {"form": form})


def mark_input_success(request):
    return render(request, "report_card/mark_input_success.html")


def home(request):
    return render(request, "report_card/home.html")


def register(request):
    if request.method == "POST":
        user_type = request.POST.get("user_type")
        if user_type == "student":
            form = StudentRegistrationForm(request.POST)
        elif user_type == "teacher":
            form = TeacherRegistrationForm(request.POST)

        if form.is_valid() and user_type == "student":
            first_name = form.cleaned_data["first_name"]
            last_name = form.cleaned_data["last_name"]
            roll_number = form.cleaned_data["roll_number"]

            # Check if the student exists in the database
            student = Student.objects.filter(
                first_name__iexact=first_name,
                last_name__iexact=last_name,
                roll_number__iexact=roll_number,
            ).first()
            print(student)
            if student:
                # Get the existing user or create a new one
                user = student.user if student.user else User()
                user.username = form.cleaned_data["username"]
                user.set_password(form.cleaned_data["password"])
                user.email = form.cleaned_data["email"]
                user.save()

                # Associate the user with the student
                student.user = user
                student.save()
                return redirect("registration_success")
            else:
                # Student not found, display an error message
                form.add_error(None, "Student not found in the database.")
                context = {"user_type": user_type, "form": form}
                return render(request, "report_card/register.html", context)

        elif form.is_valid() and user_type == "teacher":
            user = User.objects.create_user(
                first_name=form.cleaned_data["first_name"],
                last_name=form.cleaned_data["last_name"],
                username=form.cleaned_data["username"],
                password=form.cleaned_data["password"],
                email=form.cleaned_data["email"],
            )
            user.save()
            teacher = Teacher.objects.create(user=user)
            return redirect("registration_success")
        else:
            print(form.errors)
            context = {"user_type": user_type, "form": form}
            return render(request, "report_card/register.html", context)
    else:
        if request.path == "/register_as_student/":
            context = {"user_type": "student", "form": StudentRegistrationForm()}
        elif request.path == "/register_as_teacher/":
            context = {"user_type": "teacher", "form": TeacherRegistrationForm()}
        return render(request, "report_card/register.html", context)


def signin_as_student(request):
    user_type = "student"
    if request.method == "POST" and request.POST.get("user_type") == "student":
        username = request.POST.get("username")
        password = request.POST.get("password")
        user = authenticate(username=username, password=password)
        if user is not None:
            login(request, user)
            messages.success(request, "You are now logged in as a student.")
            return redirect("home")
        else:
            messages.error(request, "Invalid username or password. Try again")
    context = {"user_type": user_type}
    return render(request, "report_card/signin.html", context)


def signin_as_teacher(request):
    user_type = "teacher"
    if request.method == "POST" and request.POST.get("user_type") == "teacher":
        username = request.POST.get("username")
        password = request.POST.get("password")
        user = authenticate(username=username, password=password)
        if user is not None:
            login(request, user)
            messages.success(request, "You are now logged in as a teacher.")
            return redirect("home")
        else:
            messages.error(request, "Invalid username or password. Try again")
    context = {"user_type": user_type}
    return render(request, "report_card/signin.html", context)


def signin(request):
    # render the signin page
    if request.method == "GET":
        return render(request, "report_card/login_page.html")


def registration_success(request):
    return render(request, "report_card/registration_success.html")


@login_required
def logout_view(request):
    logout(request)
    return redirect("home")


@login_required
def generate_report_card(request):
    return render(request, "report_card/generate_report_card.html")


@login_required
def view_marks_all(request):
    # Query the database for all students
    students = Student.objects.all()

    # Create a list of dictionaries, each containing a student's details, marks, and rank
    student_marks_rank = []
    for student in students:
        marks = SubjectMarks.objects.filter(student=student)
        report_card = ReportCard.objects.filter(student=student).first()
        rank = report_card.student_rank if report_card else None
        student_marks_rank.append(
            {
                "student": student,
                "marks": marks,
                "rank": rank,
            }
        )
    # Pass the data to the template
    return render(
        request,
        "report_card/view_marks.html",
        {"student_marks_rank": student_marks_rank},
    )


@login_required
def upload_file(request):
    if request.method == "POST":
        form = UploadFileForm(request.POST, request.FILES)
        if form.is_valid():
            print("Form is valid")
            try:
                excel_file = request.FILES["file"]
                print(f"Uploaded file: {excel_file.name}")
                workbook = openpyxl.load_workbook(excel_file)
                sheet = workbook.active
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    (
                        first_name,
                        last_name,
                        roll_number,
                        math,
                        english,
                        science,
                        history,
                        comp_sci,
                    ) = row
                    print(
                        f"Creating or updating student: {first_name} {last_name} ({roll_number})"
                    )
                    Student.objects.get_or_create(
                        first_name=first_name,
                        last_name=last_name,
                        roll_number=roll_number,
                    )
                    student = Student.objects.get(roll_number=roll_number)
                    math_subject = Subject.objects.get(subject_name="Math")
                    english_subject = Subject.objects.get(subject_name="English")
                    science_subject = Subject.objects.get(subject_name="Science")
                    history_subject = Subject.objects.get(subject_name="History")
                    comp_sci_subject = Subject.objects.get(
                        subject_name="Computer Science"
                    )
                    SubjectMarks.objects.update_or_create(
                        student=student, subject=math_subject, defaults={"marks": math}
                    )
                    SubjectMarks.objects.update_or_create(
                        student=student,
                        subject=english_subject,
                        defaults={"marks": english},
                    )
                    SubjectMarks.objects.update_or_create(
                        student=student,
                        subject=science_subject,
                        defaults={"marks": science},
                    )
                    SubjectMarks.objects.update_or_create(
                        student=student,
                        subject=comp_sci_subject,
                        defaults={"marks": comp_sci},
                    )
                    SubjectMarks.objects.update_or_create(
                        student=student,
                        subject=history_subject,
                        defaults={"marks": history},
                    )

                # Calculate ranks
                for subject in Subject.objects.all():
                    marks = SubjectMarks.objects.filter(subject=subject).order_by(
                        "-marks"
                    )
                    for i, mark in enumerate(marks):
                        mark.rank = i + 1
                        mark.save()
                # Calculate overall rank
                for student in Student.objects.all():
                    total_marks = (
                        SubjectMarks.objects.filter(student=student).aggregate(
                            total=Sum("marks")
                        )["total"]
                        or 0
                    )
                    overall_rank = (
                        Student.objects.annotate(total_marks=Sum("studentmarks__marks"))
                        .filter(total_marks__gt=total_marks)
                        .count()
                        + 1
                    )
                    report_card, created = ReportCard.objects.get_or_create(
                        student=student, defaults={"student_rank": overall_rank}
                    )
                    if not created:
                        report_card.student_rank = overall_rank
                        report_card.save()

                print("File processed successfully")
                return redirect("home")
            except Exception as e:
                print(f"Error processing file: {e}")
        else:
            print("Form is invalid")
            print(form.errors)
    else:
        form = UploadFileForm()
    return render(
        request, "report_card/enter_or_update_students_details.html", {"form": form}
    )


@login_required
def view_marks(request):
    try:
        # Query the database for the current user's student record
        student = Student.objects.get(user=request.user)
    except ObjectDoesNotExist:
        # Handle the case where the student does not exist
        return render(
            request,
            "report_card/error_page.html",
            {"message": "Student does not exist"},
        )

    # Query the database for the current user's marks
    marks = SubjectMarks.objects.filter(student=student)

    # Query the database for the current user's rank
    report_card = ReportCard.objects.filter(student=student).first()
    overall_rank = report_card.student_rank if report_card else None
    # Pass the data to the template
    return render(
        request,
        "report_card/view_student_marks.html",
        {"student": student, "marks": marks, "overall_rank": overall_rank },
    )


def render_to_pdf(template_src, context_dict={}):
    template = get_template(template_src)
    html = template.render(context_dict)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
    if not pdf.err:
        return result
    return None

def view_pdf_report(request, *args, **kwargs):
    student = request.user.student
    student_marks = SubjectMarks.objects.filter(student=student)
    report_card = ReportCard.objects.get(student=student)
    pdf = render_to_pdf('report_card/generate_report_card.html', {'student': student, 'student_marks': student_marks, 'report_card': report_card})
    if pdf:
        pdf.seek(0)
        response = HttpResponse(pdf, content_type='application/pdf')
        filename = f"{student.username}_report_card.pdf"  # Customize filename if needed
        response['Content-Disposition'] = f'inline; filename="{filename}"'
        return response
    else:
        # Handle PDF generation error
        return HttpResponse("Error generating PDF", status=500)